"""Generate a Word document documenting both pipelines, tech stack, and flows.

Output: ``docs/HVAC_Cost_Estimator_Tech_Stack_and_Pipelines.docx``
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage
from PIL import ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs"
OUT_PATH = OUT_DIR / "HVAC_Cost_Estimator_Tech_Stack_and_Pipelines.docx"
ASSETS = OUT_DIR / "_docx_assets"

ACCENT = RGBColor(0x0F, 0x76, 0x6E)
GREY = RGBColor(0x47, 0x55, 0x69)
DARK = RGBColor(0x0F, 0x17, 0x2A)


def _set_run(run, *, size=11, bold=False, color=DARK, name="Calibri") -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        _set_run(run, size=16 if level == 1 else 13, bold=True, color=ACCENT)


def para(doc: Document, text: str, *, size=11, bold=False, color=DARK) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(6)


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    _set_run(run, size=11)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        _set_run(run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        tc_pr = hdr[i]._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "0F766E")
        shading.set(qn("w:val"), "clear")
        tc_pr.append(shading)
    for r_i, row in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        for c_i, val in enumerate(row):
            cells[c_i].text = ""
            p = cells[c_i].paragraphs[0]
            run = p.add_run(val)
            _set_run(run, size=9, color=DARK)
    doc.add_paragraph()


def _font(size: int):
    for name in (
        "C:\\Windows\\Fonts\\segoeui.ttf",
        "C:\\Windows\\Fonts\\arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def _draw_flowchart(
    path: Path,
    title: str,
    boxes: list[tuple[str, str]],
    *,
    width: int = 1100,
) -> None:
    """Draw a vertical flowchart: list of (label, subtitle)."""
    box_h = 58
    gap = 28
    top = 70
    height = top + len(boxes) * (box_h + gap) + 40
    img = PILImage.new("RGB", (width, height), (248, 250, 252))
    draw = ImageDraw.Draw(img)
    title_font = _font(22)
    label_font = _font(15)
    sub_font = _font(12)

    draw.text((40, 22), title, fill=(15, 118, 110), font=title_font)

    x0, x1 = 80, width - 80
    cx = width // 2
    for i, (label, sub) in enumerate(boxes):
        y0 = top + i * (box_h + gap)
        y1 = y0 + box_h
        # box
        draw.rounded_rectangle(
            [x0, y0, x1, y1],
            radius=10,
            fill=(15, 118, 110) if i == 0 or i == len(boxes) - 1 else (255, 255, 255),
            outline=(15, 118, 110),
            width=2,
        )
        fill = (255, 255, 255) if i == 0 or i == len(boxes) - 1 else (15, 23, 42)
        sub_fill = (226, 232, 240) if i == 0 or i == len(boxes) - 1 else (71, 85, 105)
        # text
        draw.text((x0 + 18, y0 + 10), label, fill=fill, font=label_font)
        draw.text((x0 + 18, y0 + 32), sub, fill=sub_fill, font=sub_font)
        # arrow to next
        if i < len(boxes) - 1:
            ay0 = y1 + 2
            ay1 = y1 + gap - 6
            draw.line([(cx, ay0), (cx, ay1)], fill=(15, 118, 110), width=3)
            draw.polygon(
                [(cx - 7, ay1 - 10), (cx + 7, ay1 - 10), (cx, ay1)],
                fill=(15, 118, 110),
            )

    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def _draw_parallel_overview(path: Path) -> None:
    """Top-level diagram: one PDF fans into three tracks then merges."""
    w, h = 1200, 720
    img = PILImage.new("RGB", (w, h), (248, 250, 252))
    draw = ImageDraw.Draw(img)
    title_f = _font(22)
    box_f = _font(14)
    small_f = _font(11)
    draw.text((40, 24), "End-to-End System Flow (on PDF upload)", fill=(15, 118, 110), font=title_f)

    def box(xy, text, *, fill=(255, 255, 255), outline=(15, 118, 110), tw=None):
        x0, y0, x1, y1 = xy
        draw.rounded_rectangle([x0, y0, x1, y1], radius=8, fill=fill, outline=outline, width=2)
        color = (255, 255, 255) if fill == (15, 118, 110) else (15, 23, 42)
        draw.text((x0 + 12, y0 + 14), text, fill=color, font=box_f if tw is None else small_f)

    def arrow(a, b):
        draw.line([a, b], fill=(15, 118, 110), width=3)
        bx, by = b
        draw.polygon([(bx - 8, by - 6), (bx - 8, by + 6), (bx, by)], fill=(15, 118, 110))

    # Upload
    box((480, 70, 720, 120), "PDF Upload (React + FastAPI)", fill=(15, 118, 110))

    # Three parallel tracks
    tracks = [
        (60, "Track A — Requirements", "PyMuPDF text + heuristics\n→ requirement.pdf"),
        (420, "Track B — Tech Symbols", "Spatial legend parse + crops\n→ technical symbol.pdf"),
        (780, "Track C — CV Costing", "Render → ROI/OCR + detect\n→ costed report"),
    ]
    for x, title, body in tracks:
        box((x, 200, x + 340, 250), title, fill=(15, 118, 110))
        box((x, 270, x + 340, 360), body, tw=True)
        arrow((600, 120), (x + 170, 200))

    # Merge
    box((360, 440, 840, 500), "SQLite project record + downloadable artifacts", fill=(15, 118, 110))
    for x, _, _ in tracks:
        arrow((x + 170, 360), (600, 440))

    box((360, 560, 840, 640), "React Dashboard\nMetadata · Cost grid · CSV · Requirement PDF · Symbol PDF")
    arrow((600, 500), (600, 560))

    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def build() -> Path:
    ASSETS.mkdir(parents=True, exist_ok=True)
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    overview_img = ASSETS / "overview_flow.png"
    req_img = ASSETS / "requirement_flow.png"
    sym_img = ASSETS / "symbol_flow.png"
    cv_img = ASSETS / "cv_costing_flow.png"

    _draw_parallel_overview(overview_img)
    _draw_flowchart(
        req_img,
        "Pipeline A — Requirement Extraction",
        [
            ("1. Open full PDF", "PyMuPDF — all pages, no CV page cap"),
            ("2. Detect sheet family", "Garland AGENT: vs Architect PROJECT OWNER"),
            ("3. Extract provider / project fields", "Agent, company, phone, address, title"),
            ("4. Parse scope of work", "Numbered items + trade SCOPE sections"),
            ("5. Write requirement PDF", "ReportLab → storage/requirements/"),
        ],
    )
    _draw_flowchart(
        sym_img,
        "Pipeline B — Technical Symbol Legend Extraction",
        [
            ("1. Scan every page for legend headings", "SYMBOL LEGEND / TECHNOLOGY SYMBOL LEGEND"),
            ("2a. Technology legend (spatial)", "Rotated X-rows / Y-columns; pack MFG/Part/Notes"),
            ("2b. Numbered Garland legend", "Parse 1…N descriptions; crop glyph left of N."),
            ("3. Rotation-correct pixmap crop", "page.rotation_matrix + grid-border strip"),
            ("4. Write technical symbol PDF", "ReportLab embeds glyph + number when both exist"),
        ],
    )
    _draw_flowchart(
        cv_img,
        "Pipeline C — CV Cost Estimation (device counts)",
        [
            ("1. PDF → page images", "pdf2image / PyMuPDF @ 300 DPI (page-capped)"),
            ("2. Text branch", "ROI detect → PaddleOCR → rapidfuzz metadata map"),
            ("3. Symbol branch", "Symbol detect → ONNX device classify → counts"),
            ("4. Cost calculation", "counts × rate table → line items + grand total"),
            ("5. Persist + UI", "SQLite → React editable grid + CSV export"),
        ],
    )

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # Cover
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Mechanical Device Cost Estimator")
    _set_run(r, size=26, bold=True, color=ACCENT)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("Tech Stack, Pipelines & Process Flows")
    _set_run(r, size=16, color=GREY)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "i-TAB HVAC POC  ·  Local-only  ·  FastAPI + React + SQLite\n"
        "Documents both text-extraction pipelines and the CV costing pipeline"
    )
    _set_run(r, size=10, color=GREY)

    doc.add_paragraph()

    # 1. Overview
    heading(doc, "1. System overview", 1)
    para(
        doc,
        "On each PDF upload the backend runs three complementary tracks in parallel. "
        "Tracks A and B are text/geometry extractors over the full PDF (no ML page cap). "
        "Track C is the computer-vision costing pipeline over rendered page images "
        "(subject to max_pdf_pages). Results land in SQLite and as downloadable PDFs.",
    )
    doc.add_picture(str(overview_img), width=Inches(6.5))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run("Figure 1 — End-to-end system flow")
    _set_run(r, size=9, color=GREY)

    # 2. Tech stack
    heading(doc, "2. Shared tech stack", 1)
    add_table(
        doc,
        ["Layer", "Technology", "Role"],
        [
            ["API", "Python 3.10+, FastAPI, Uvicorn", "Upload, project CRUD, background jobs"],
            ["Config", "pydantic-settings", "Env-driven Settings (mock vs real models)"],
            ["Database", "SQLite + SQLAlchemy 2.x", "Projects, pages, device lines, paths"],
            ["PDF I/O", "PyMuPDF (fitz), ReportLab, Pillow", "Text extract, render, output PDFs"],
            ["PDF render", "pdf2image (+ Poppler) / PyMuPDF fallback", "300 DPI page PNGs for CV"],
            ["Fuzzy match", "rapidfuzz", "Title-block label → field aliases"],
            ["Frontend", "React 18, TypeScript, Vite, Tailwind", "Dashboard, upload, editable grid"],
            ["Testing", "pytest, httpx, Vitest", "Backend + frontend unit/API tests"],
            ["Hosting", "Local only (no Docker / cloud)", "Single-user POC"],
        ],
    )

    heading(doc, "2.1 Optional ML stack (real models)", 2)
    para(
        doc,
        "Default HVAC_USE_MOCK_MODELS=true uses deterministic mock detectors so the "
        "POC runs without heavy weights. Real backends (requirements-ml.txt) plug in "
        "behind the same Protocols:",
    )
    add_table(
        doc,
        ["Component", "Library", "Notes"],
        [
            ["ROI / symbol detect", "Detectron2 Faster R-CNN", "No Windows wheels by default"],
            ["OCR", "PaddleOCR (en)", "Title-block text lines"],
            ["Device classify", "timm CNN → ONNX Runtime", "Crop → device type"],
        ],
    )

    # 3. Pipeline A
    heading(doc, "3. Pipeline A — Requirement extraction", 1)
    para(
        doc,
        "Purpose: pull provider/contact metadata and scope-of-work from bid/drawing "
        "sets into a readable ``<stem> requirement.pdf``.",
    )
    heading(doc, "3.1 Tech stack", 2)
    add_table(
        doc,
        ["Piece", "Module / tool", "Detail"],
        [
            ["Extract", "ml/requirement_extractor.py", "PyMuPDF text + family heuristics"],
            ["Families", "Garland / Architect CD", "AGENT: block vs PROJECT OWNER cover"],
            ["Scope", "Numbered + trade SCOPE parsers", "Handles split ``1.`` / body lines"],
            ["Output PDF", "ml/requirement_pdf.py + ReportLab", "storage/requirements/"],
            ["API", "GET …/requirement.pdf", "Download from dashboard"],
            ["CLI", "scripts/extract_requirements.py", "Batch folder processing"],
        ],
    )
    heading(doc, "3.2 Process", 2)
    for step in [
        "Open the PDF with PyMuPDF and read embedded text (full page count).",
        "Classify layout family (Garland vs Architect CD) from headings/markers.",
        "Map agent, company, phones, address, project title, sheet indexes.",
        "Collect scope sections (numbered lists and/or trade SCOPE bodies).",
        "Generate requirement PDF; store path on the Project row.",
    ]:
        bullet(doc, step)
    doc.add_picture(str(req_img), width=Inches(5.8))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run("Figure 2 — Requirement extraction flow")
    _set_run(r, size=9, color=GREY)

    # 4. Pipeline B
    heading(doc, "4. Pipeline B — Technical symbol legend extraction", 1)
    para(
        doc,
        "Purpose: extract TECHNOLOGY SYMBOL LEGEND and numbered SYMBOL LEGEND "
        "tables with real CAD glyphs (not ML classification) into "
        "``<stem> technical symbol.pdf``.",
    )
    heading(doc, "4.1 Tech stack", 2)
    add_table(
        doc,
        ["Piece", "Module / tool", "Detail"],
        [
            ["Extract", "ml/symbol_table_extractor.py", "Text spans + spatial clustering"],
            ["Tech legend", "X-row / Y-column spatial parse", "Handles page.rotation 90/270"],
            ["Numbered legend", "Line parse + glyph crop left of N.", "Blank when sheet has no icon"],
            ["Crop / clean", "PyMuPDF pixmap + Pillow", "rotation_matrix; strip table borders"],
            ["Fallback icons", "Pillow draw helpers", "Only if tech-legend crop is blank"],
            ["Output PDF", "ml/symbol_table_pdf.py + ReportLab", "Glyph + number side-by-side"],
            ["API", "GET …/technical-symbol.pdf", "Download from dashboard"],
            ["CLI", "scripts/extract_technical_symbols.py", "Batch / single PDF"],
        ],
    )
    heading(doc, "4.2 Process", 2)
    for step in [
        "Scan every page for SYMBOL LEGEND / TECHNOLOGY SYMBOL LEGEND headings.",
        "Technology path: locate column headers, cluster description rows, pack MFG/Part/Notes.",
        "Crop Symbol column cells with rotation-corrected visual rectangles; strip grid L-brackets.",
        "Numbered path: parse 1…N descriptions; crop ink immediately left of each ``N.`` label.",
        "When both glyph and number exist, PDF Symbol cell shows [glyph] [N.] like the drawing.",
        "Write technical symbol PDF under storage/technical_symbols/.",
    ]:
        bullet(doc, step)
    doc.add_picture(str(sym_img), width=Inches(5.8))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run("Figure 3 — Technical symbol extraction flow")
    _set_run(r, size=9, color=GREY)

    # 5. Pipeline C
    heading(doc, "5. Pipeline C — CV cost estimation", 1)
    para(
        doc,
        "Purpose: from rendered pages, extract title-block metadata and count HVAC "
        "device symbols, then join to a unit-cost rate table for an editable report.",
    )
    heading(doc, "5.1 Tech stack", 2)
    add_table(
        doc,
        ["Stage", "Module", "Mock / Real"],
        [
            ["PDF → images", "ml/pdf_to_image.py", "Always real (pdf2image / PyMuPDF)"],
            ["Title-block ROI", "ml/roi_detector.py", "Mock strip / Detectron2"],
            ["OCR", "ml/ocr.py", "Mock lines / PaddleOCR"],
            ["Metadata map", "ml/metadata_mapper.py", "Always real (rapidfuzz)"],
            ["Symbol detect", "ml/symbol_detector.py", "Mock boxes / Detectron2"],
            ["Classify", "ml/classifier.py", "Mock hash / ONNX CNN"],
            ["Costing", "ml/cost_calculator.py", "Always real (rate table join)"],
            ["Orchestration", "ml/pipeline.py", "ThreadPool + background task"],
        ],
    )
    heading(doc, "5.2 Process", 2)
    for step in [
        "Rasterize PDF pages to PNGs (DPI and page cap from Settings).",
        "Text branch: detect title-block ROI → OCR lines → fuzzy map to schema fields.",
        "Symbol branch: detect device boxes → classify type → aggregate counts.",
        "Join counts to rate table; compute line totals and grand total server-side.",
        "Persist Project / Page / DeviceLine; UI polls until status=done; CSV export available.",
    ]:
        bullet(doc, step)
    doc.add_picture(str(cv_img), width=Inches(5.8))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run("Figure 4 — CV costing pipeline flow")
    _set_run(r, size=9, color=GREY)

    # 6. Orchestration
    heading(doc, "6. Orchestration & API", 1)
    para(
        doc,
        "run_pipeline() starts requirement extract, technical-symbol extract, and "
        "page render concurrently (ThreadPoolExecutor, 3 workers), then runs the CV "
        "stages. process_project() wraps this as a FastAPI BackgroundTasks job with "
        "status pending → processing → done | failed.",
    )
    add_table(
        doc,
        ["Endpoint", "Purpose"],
        [
            ["POST /api/upload", "Validate PDF, create project, start pipeline"],
            ["GET /api/projects", "List projects + status"],
            ["GET /api/projects/{id}", "Detail: metadata, lines, totals, PDF flags"],
            ["PATCH /api/projects/{id}/lines/{line_id}", "Override count / unit cost"],
            ["GET /api/projects/{id}/report", "JSON report"],
            ["GET /api/projects/{id}/report/csv", "CSV export"],
            ["GET /api/projects/{id}/requirement.pdf", "Download requirement extract"],
            ["GET /api/projects/{id}/technical-symbol.pdf", "Download symbol legend extract"],
        ],
    )

    # 7. Artifacts
    heading(doc, "7. Output artifacts", 1)
    add_table(
        doc,
        ["Artifact", "Location", "Producer"],
        [
            ["Requirement PDF", "backend/storage/requirements/", "Pipeline A"],
            ["Technical symbol PDF", "backend/storage/technical_symbols/", "Pipeline B"],
            ["Page PNGs", "backend/storage/…/pages/", "Pipeline C render"],
            ["Cost report", "SQLite + API/CSV", "Pipeline C costing"],
        ],
    )

    # 8. Lead talking points
    heading(doc, "8. Lead talking points", 1)
    for tip in [
        "Three tracks from one upload: requirements, legend symbols, and device costing.",
        "Legend glyphs are cropped from the PDF geometry (rotation-aware) — not invented by ML.",
        "Numbered SYMBOL LEGEND shows [glyph] [N.] when both exist on the sheet; blank cells stay blank.",
        "CV models are swappable (mock ↔ Detectron2/Paddle/ONNX) behind stable Protocols.",
        "POC is local-only: FastAPI + React + SQLite; no containers or paid APIs required.",
    ]:
        bullet(doc, tip)

    footer = doc.add_paragraph()
    r = footer.add_run(
        "Generated for the i-TAB HVAC Cost Estimator POC. "
        "Source: hvac-cost-estimator/docs + backend/ml/pipeline.py."
    )
    _set_run(r, size=9, color=GREY)

    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
    sys.exit(0)
